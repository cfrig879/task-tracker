from pathlib import Path
import calendar
import openpyxl
import re
from datetime import date, datetime, timedelta

from django.conf import settings
from django.contrib import messages
from django.db.models import Count, Q
from django.http import FileResponse, Http404
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from django.contrib.auth import login
from django.contrib.auth.decorators import login_required

from .category_forms import CategoryForm
from .import_forms import TaskImportForm
from .models import Category, Task

from .forms import TaskForm, SignUpForm

def landing(request):
    return render(request, "tracker/landing.html")

def demo_dashboard(request):
    return render(request, "tracker/demo_dashboard.html")

def signup(request):
    if request.user.is_authenticated:
        return redirect("dashboard")

    if request.method == "POST":
        form = SignUpForm(request.POST)
        if form.is_valid():
            user = form.save()
            login(request, user)
            messages.success(request, "Your account has been created.")
            return redirect("dashboard")
    else:
        form = SignUpForm()

    return render(request, "tracker/signup.html", {"form": form})


def download_import_template(request):
    file_path = Path(settings.BASE_DIR) / "tracker" / "static" / "tracker" / "task_import_template.xlsx"

    if not file_path.exists():
        raise Http404("Import template not found.")

    return FileResponse(
        open(file_path, "rb"),
        as_attachment=True,
        filename="task_import_template.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _format_minutes(total_minutes: int | None) -> str:
    if not total_minutes:
        return "0m"
    hours = total_minutes // 60
    minutes = total_minutes % 60
    if hours and minutes:
        return f"{hours}h {minutes}m"
    if hours:
        return f"{hours}h"
    return f"{minutes}m"


def _parse_bool(value: str) -> bool:
    return (value or "").strip().lower() in {"1", "true", "yes", "y", "t"}

def _normalize_category_name(raw: str) -> str:
    cleaned = re.sub(r"\s+", " ", (raw or "").strip())
    if not cleaned:
        return ""
    return cleaned.title()

def _parse_csv_date(value: str) -> date | None:
    raw = (value or "").strip()
    if not raw:
        return None

    formats = (
        "%Y-%m-%d",
        "%Y-%m-%d %H:%M:%S",
        "%m/%d/%Y",
        "%m/%d/%y",
    )

    for fmt in formats:
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue

    raise ValueError(
        f"invalid date '{raw}' (use YYYY-MM-DD or M/D/YYYY)"
    )


def _parse_csv_datetime(value: str) -> datetime | None:
    raw = (value or "").strip()
    if not raw:
        return None

    formats = (
        "%Y-%m-%d %H:%M",
        "%Y-%m-%dT%H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%m/%d/%Y %H:%M",
        "%m/%d/%Y %I:%M",
        "%m/%d/%Y %I:%M %p",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %I:%M:%S %p",
        "%m/%d/%y %H:%M",
        "%m/%d/%y %I:%M",
        "%m/%d/%y %I:%M %p",
        "%m/%d/%y %H:%M:%S",
        "%m/%d/%y %I:%M:%S %p",
    )

    parsed: datetime | None = None
    for fmt in formats:
        try:
            parsed = datetime.strptime(raw, fmt)
            break
        except ValueError:
            continue

    if parsed is None:
        raise ValueError(
            f"invalid datetime '{raw}' "
            f"(use YYYY-MM-DD HH:MM, YYYY-MM-DDTHH:MM, or spreadsheet-style dates)"
        )

    if settings.USE_TZ and timezone.is_naive(parsed):
        parsed = timezone.make_aware(parsed, timezone.get_current_timezone())

    return parsed


def _parse_estimated_minutes(row: dict[str, str]) -> int | None:
    hours_raw = (row.get("estimated_hours") or "").strip()
    minutes_raw = (row.get("estimated_minutes_part") or "").strip()
    total_raw = (row.get("estimated_minutes") or "").strip()

    if total_raw:
        total = int(total_raw)
        return total if total > 0 else None

    hours = int(hours_raw) if hours_raw else 0
    minutes = int(minutes_raw) if minutes_raw else 0
    total = hours * 60 + minutes
    return total if total > 0 else None

@login_required
def dashboard(request):
    now = timezone.localtime()
    today = now.date()
    week_start = today - timedelta(days=today.weekday())
    week_end = week_start + timedelta(days=6)
    next_7_days = today + timedelta(days=7)

    base_qs = Task.objects.filter(user=request.user).select_related("category")

    in_progress_tasks = list(
        base_qs.filter(completed=False, start_at__isnull=False, end_at__isnull=True)
        .order_by("start_at", "due_date", "title")[:8]
    )

    today_tasks = list(
        base_qs.filter(completed=False, due_date=today)
        .order_by("start_at", "title")[:8]
    )

    overdue_tasks = list(
        base_qs.filter(completed=False, due_date__lt=today)
        .order_by("due_date", "title")[:12]
    )

    upcoming_tasks = list(
        base_qs.filter(completed=False, due_date__gte=today, due_date__lte=next_7_days)
        .order_by("due_date", "start_at", "title")[:12]
    )

    overdue_count = base_qs.filter(completed=False, due_date__lt=today).count()
    due_today_count = base_qs.filter(completed=False, due_date=today).count()
    missing_estimate_count = base_qs.filter(
        completed=False,
        estimated_minutes__isnull=True,
    ).count()

    missing_actual_duration_count = base_qs.filter(completed=True).filter(
        Q(start_at__isnull=True) | Q(end_at__isnull=True)
    ).count()

    uncategorized_count = base_qs.filter(completed=False, category__isnull=True).count()

    week_tasks = list(base_qs.filter(due_date__range=(week_start, week_end)))

    completed_this_week_count = sum(1 for task in week_tasks if task.completed)

    comparable_week_tasks = [
        task
        for task in week_tasks
        if task.completed
        and task.estimated_minutes is not None
        and task.duration_minutes is not None
    ]

    planned_this_week_minutes = sum(task.estimated_minutes for task in comparable_week_tasks)
    actual_this_week_minutes = sum(task.duration_minutes for task in comparable_week_tasks)

    planned_this_week_display = _format_minutes(planned_this_week_minutes)
    actual_this_week_display = _format_minutes(actual_this_week_minutes)

    underestimated_count = sum(
        1 for task in comparable_week_tasks
        if task.duration_minutes > task.estimated_minutes
    )

    overestimated_count = sum(
        1 for task in comparable_week_tasks
        if task.duration_minutes < task.estimated_minutes
    )

    missing_estimate_this_week_count = sum(
        1 for task in week_tasks
        if task.estimated_minutes is None
    )

    completed_without_duration_count = sum(
        1 for task in week_tasks
        if task.completed and task.duration_minutes is None
    )

    largest_overrun_category_name = "None yet"
    overrun_by_category: dict[str, int] = {}

    for task in comparable_week_tasks:
        diff = task.duration_minutes - task.estimated_minutes
        if diff > 0:
            category_name = task.category.name if task.category else "Uncategorized"
            overrun_by_category[category_name] = overrun_by_category.get(category_name, 0) + diff

    if overrun_by_category:
        largest_overrun_category_name = max(overrun_by_category, key=overrun_by_category.get)

    total_recent_entries = len(week_tasks)
    sufficiently_detailed_entries = len(comparable_week_tasks)

    insight_completion_percent = (
        round((sufficiently_detailed_entries / total_recent_entries) * 100)
        if total_recent_entries
        else 0
    )

    weekly_signals: list[str] = []

    if comparable_week_tasks:
        if actual_this_week_minutes > planned_this_week_minutes:
            weekly_signals.append(
                "Completed tasks with timing data took longer than planned this week."
            )
        elif actual_this_week_minutes < planned_this_week_minutes:
            weekly_signals.append(
                "Completed tasks with timing data took less time than planned this week."
            )
        else:
            weekly_signals.append(
                "Completed tasks with timing data matched the planned time this week."
            )

        if underestimated_count:
            weekly_signals.append(
                f"{underestimated_count} completed task{'s' if underestimated_count != 1 else ''} took longer than estimated."
            )

        if overestimated_count:
            weekly_signals.append(
                f"{overestimated_count} completed task{'s' if overestimated_count != 1 else ''} took less time than estimated."
            )
    else:
        weekly_signals.append(
            "Add estimates and start/end times to completed tasks to unlock stronger weekly insights."
        )

    if completed_without_duration_count:
        weekly_signals.append(
            f"{completed_without_duration_count} completed task{'s' if completed_without_duration_count != 1 else ''} still need actual timing data."
        )

    if missing_estimate_this_week_count:
        weekly_signals.append(
            f"{missing_estimate_this_week_count} task{'s' if missing_estimate_this_week_count != 1 else ''} this week are missing estimates."
        )

    context = {
        "in_progress_tasks": in_progress_tasks,
        "today_tasks": today_tasks,
        "overdue_tasks": overdue_tasks,
        "upcoming_tasks": upcoming_tasks,

        "overdue_count": overdue_count,
        "due_today_count": due_today_count,
        "missing_estimate_count": missing_estimate_count,
        "missing_actual_duration_count": missing_actual_duration_count,
        "uncategorized_count": uncategorized_count,

        "completed_this_week_count": completed_this_week_count,
        "planned_this_week_display": planned_this_week_display,
        "actual_this_week_display": actual_this_week_display,
        "largest_overrun_category_name": largest_overrun_category_name,

        "weekly_signals": weekly_signals,
        "insight_completion_percent": insight_completion_percent,
    }

    return render(request, "tracker/dashboard.html", context)

@login_required
def add_task(request):
    if request.method == "POST":
        form = TaskForm(request.POST, user=request.user)
        if form.is_valid():
            task = form.save(commit=False)
            task.user = request.user
            task.save()
            form.save_m2m()
            return redirect(task)
    else:
        form = TaskForm(user=request.user)

    return render(request, "tracker/add_task.html", {"form": form})

@login_required
def task_detail(request, pk: int):
    task = get_object_or_404(
        Task.objects.filter(user=request.user).select_related("category"),
        pk=pk,
    )
    return render(request, "tracker/task_detail.html", {"task": task})

@login_required
def edit_task(request, pk: int):
    task = get_object_or_404(Task, pk=pk, user=request.user)

    if request.method == "POST":
        form = TaskForm(request.POST, instance=task, user=request.user)
        if form.is_valid():
            task = form.save(commit=False)
            task.user = request.user
            task.save()
            form.save_m2m()
            return redirect(task)
    else:
        form = TaskForm(instance=task, user=request.user)

    return render(request, "tracker/edit_task.html", {"form": form, "task": task})

@login_required
def delete_task(request, pk: int):
    task = get_object_or_404(Task, pk=pk, user=request.user)

    if request.method == "POST":
        task.delete()
        messages.success(request, "Task deleted.")
        return redirect("task_list")

    return render(request, "tracker/delete_task.html", {"task": task})

@login_required
@require_POST
def toggle_task_complete(request, pk: int):
    task = get_object_or_404(Task, pk=pk, user=request.user)
    task.completed = not task.completed
    task.save(update_fields=["completed"])
    return redirect(task)

@login_required
def categories(request):
    if request.method == "POST":
        form = CategoryForm(request.POST)
        if form.is_valid():
            category = form.save(commit=False)
            category.user = request.user
            category.save()
            messages.success(request, "Category added.")
            return redirect("categories")
    else:
        form = CategoryForm()

    categories_qs = (
        Category.objects.filter(user=request.user)
        .annotate(task_count=Count("tasks"))
        .order_by("name")
    )
    uncategorized_count = Task.objects.filter(
        user=request.user,
        category__isnull=True,
    ).count()

    context = {
        "form": form,
        "categories": categories_qs,
        "uncategorized_count": uncategorized_count,
    }
    return render(request, "tracker/categories.html", context)


@login_required
def edit_category(request, pk: int):
    category = get_object_or_404(Category, pk=pk, user=request.user)

    if request.method == "POST":
        form = CategoryForm(request.POST, instance=category)
        if form.is_valid():
            category = form.save(commit=False)
            category.user = request.user
            category.save()
            messages.success(request, "Category updated.")
            return redirect("categories")
    else:
        form = CategoryForm(instance=category)

    context = {
        "form": form,
        "category": category,
    }
    return render(request, "tracker/edit_category.html", context)


@login_required
def delete_category(request, pk: int):
    category = get_object_or_404(Category, pk=pk, user=request.user)

    if request.method == "POST":
        category.delete()
        messages.success(request, "Category deleted.")
        return redirect("categories")

    context = {
        "category": category,
    }
    return render(request, "tracker/delete_category.html", context)


def about(request):
    return render(request, "tracker/about.html")

@login_required
def calendar_view(request):
    today = date.today()

    if request.GET.get("today") == "1":
        year = today.year
        month = today.month
    else:
        year = int(request.GET.get("year", today.year))
        month = int(request.GET.get("month", today.month))

    cal = calendar.Calendar(firstweekday=6)
    weeks = cal.monthdatescalendar(year, month)

    grid_start = weeks[0][0]
    grid_end = weeks[-1][-1]

    tasks = (
        Task.objects.filter(
            user=request.user,
            due_date__isnull=False,
            due_date__range=(grid_start, grid_end),
        )
        .select_related("category")
        .order_by("due_date", "completed", "title")
    )

    tasks_by_day: dict[date, list[Task]] = {}
    for task in tasks:
        tasks_by_day.setdefault(task.due_date, []).append(task)

    prev_month = month - 1
    prev_year = year
    if prev_month == 0:
        prev_month = 12
        prev_year -= 1

    next_month = month + 1
    next_year = year
    if next_month == 13:
        next_month = 1
        next_year += 1

    context = {
        "year": year,
        "month": month,
        "month_name": calendar.month_name[month],
        "weeks": weeks,
        "tasks_by_day": tasks_by_day,
        "today": today,
        "prev_year": prev_year,
        "prev_month": prev_month,
        "next_year": next_year,
        "next_month": next_month,
    }
    return render(request, "tracker/calendar.html", context)

@login_required
def import_tasks(request):
    if request.method == "POST":
        form = TaskImportForm(request.POST, request.FILES)
        if form.is_valid():
            file_obj = form.cleaned_data["file"]

            if not file_obj.name.lower().endswith(".xlsx"):
                form.add_error("file", "Please upload the Cadence Excel template (.xlsx).")
            else:
                created = 0
                skipped = 0
                errors: list[str] = []

                try:
                    wb = openpyxl.load_workbook(file_obj, data_only=True)
                except Exception:
                    form.add_error(
                        "file",
                        "That file could not be read as an Excel workbook. Please use the Cadence template.",
                    )
                    wb = None

                if wb is not None:
                    if "Tasks" not in wb.sheetnames:
                        form.add_error(
                            "file",
                            "Invalid template: missing the 'Tasks' sheet. Please download the Cadence template and try again.",
                        )
                    else:
                        ws = wb["Tasks"]

                        # Find header row containing 'title'
                        header_row_idx = None
                        header_map: dict[str, int] = {}
                        max_cols = min(ws.max_column, 40)

                        for r in range(1, min(ws.max_row, 25) + 1):
                            values = [ws.cell(r, c).value for c in range(1, max_cols + 1)]
                            normalized = [(str(v).strip().lower() if v is not None else "") for v in values]
                            if "title" in normalized:
                                header_row_idx = r
                                for c, name in enumerate(normalized, start=1):
                                    if name:
                                        header_map[name] = c
                                break

                        if header_row_idx is None:
                            form.add_error(
                                "file",
                                "Invalid template: could not find a header row with a 'title' column.",
                            )
                        else:
                            required_cols = {
                                "title",
                                "category",
                                "due_date",
                                "start_at",
                                "end_at",
                                "estimated_hours",
                                "estimated_minutes_part",
                                "notes",
                                "completed",
                            }
                            missing = required_cols - set(header_map.keys())
                            if missing:
                                form.add_error(
                                    "file",
                                    f"Invalid template: missing column(s): {', '.join(sorted(missing))}.",
                                )
                            else:
                                def cell_str(row_num: int, col_name: str) -> str:
                                    val = ws.cell(row_num, header_map[col_name]).value
                                    return "" if val is None else str(val).strip()

                                for row_num in range(header_row_idx + 1, ws.max_row + 1):
                                    title = cell_str(row_num, "title")
                                    if not title:
                                        # skip completely blank rows
                                        row_values = [ws.cell(row_num, c).value for c in header_map.values()]
                                        if all(v is None or (isinstance(v, str) and not v.strip()) for v in row_values):
                                            continue
                                        skipped += 1
                                        errors.append(f"Row {row_num}: title is required.")
                                        continue

                                    cat_name_raw = cell_str(row_num, "category")
                                    cat_name = _normalize_category_name(cat_name_raw)

                                    category = None
                                    if cat_name:
                                        category, _ = Category.objects.get_or_create(
                                            user=request.user,
                                            name=cat_name,
                                        )

                                    try:
                                        due_date = _parse_csv_date(cell_str(row_num, "due_date"))
                                    except ValueError:
                                        skipped += 1
                                        errors.append(
                                            f"Row {row_num}: invalid due_date '{cell_str(row_num, 'due_date')}' (use YYYY-MM-DD)."
                                        )
                                        continue

                                    try:
                                        start_at = _parse_csv_datetime(cell_str(row_num, "start_at"))
                                    except ValueError as exc:
                                        skipped += 1
                                        errors.append(f"Row {row_num}: {exc}.")
                                        continue

                                    try:
                                        end_at = _parse_csv_datetime(cell_str(row_num, "end_at"))
                                    except ValueError as exc:
                                        skipped += 1
                                        errors.append(f"Row {row_num}: {exc}.")
                                        continue

                                    if start_at and end_at and end_at < start_at:
                                        skipped += 1
                                        errors.append(f"Row {row_num}: end_at must be after start_at.")
                                        continue

                                    row_dict = {
                                        "estimated_hours": cell_str(row_num, "estimated_hours"),
                                        "estimated_minutes_part": cell_str(row_num, "estimated_minutes_part"),
                                        "estimated_minutes": "",
                                    }

                                    try:
                                        estimated_minutes = _parse_estimated_minutes(row_dict)
                                    except ValueError:
                                        skipped += 1
                                        errors.append(
                                            f"Row {row_num}: invalid estimate values (estimated_hours and estimated_minutes_part must be integers)."
                                        )
                                        continue

                                    notes = cell_str(row_num, "notes")
                                    completed = _parse_bool(cell_str(row_num, "completed"))

                                    Task.objects.create(
                                        user=request.user,
                                        title=title,
                                        category=category,
                                        due_date=due_date,
                                        start_at=start_at,
                                        end_at=end_at,
                                        estimated_minutes=estimated_minutes,
                                        notes=notes,
                                        completed=completed,
                                    )
                                    created += 1

                                messages.success(
                                    request,
                                    f"Imported {created} task(s). Skipped {skipped}.",
                                )
                                if errors:
                                    messages.warning(
                                        request,
                                        "Some rows were skipped. Review errors on the import page.",
                                    )

                                request.session["import_errors"] = errors[:50]
                                return redirect("import_tasks")
    else:
        form = TaskImportForm()

    errors = request.session.pop("import_errors", [])
    return render(
        request,
        "tracker/import_tasks.html",
        {"form": form, "errors": errors},
    )

@login_required
def task_list(request):
    tasks = Task.objects.filter(user=request.user).select_related("category")

    search = (request.GET.get("search") or "").strip()
    status = (request.GET.get("status") or "").strip()
    category_id = (request.GET.get("category") or "").strip()
    sort = (request.GET.get("sort") or "").strip()

    if search:
        tasks = tasks.filter(title__icontains=search)

    if status == "open":
        tasks = tasks.filter(completed=False)
    elif status == "completed":
        tasks = tasks.filter(completed=True)

    if category_id:
        tasks = tasks.filter(category_id=category_id)

    sort_map = {
        "title": "title",
        "-title": "-title",
        "due_date": "due_date",
        "-due_date": "-due_date",
        "updated_at": "updated_at",
        "-updated_at": "-updated_at",
        "created_at": "created_at",
        "-created_at": "-created_at",
        "status": "completed",
        "-status": "-completed",
    }

    tasks = tasks.order_by(sort_map.get(sort, "-updated_at"), "title")

    categories = Category.objects.filter(user=request.user).order_by("name")

    context = {
        "tasks": tasks,
        "categories": categories,
        "search": search,
        "status": status,
        "category_id": category_id,
        "sort": sort,
    }
    return render(request, "tracker/task_list.html", context)